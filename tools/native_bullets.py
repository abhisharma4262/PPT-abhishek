"""
native_bullets: makes every list one text box with native PowerPoint bullets.

The owner wants each list in a single text box, so it can be moved, aligned and edited as a block, with
bullets that belong to the text (no separate dot shapes to realign). Bullet style (chosen Oct 2026):
a teal dot (Pulse on dark panels), Arial U+25CF at 86% of the text size, hanging indent; check-mark lists
use the Wingdings check in teal. Bold key phrases stay as bold runs.

  python3 native_bullets.py in.pptx out.pptx              # lists tagged by mkit (html2pptx runs this for you)
  python3 native_bullets.py in.pptx out.pptx --cards 2,5  # existing slides: dots + text boxes inside cards

Tag mode: mkit's dot_list() / tick_list() tag each item ("l:" in the shape name, added by extract.js);
every tagged list becomes one text box at the list's position, with the item spacing and line height
from the HTML. Card mode is for decks already built or edited by hand: inside each rounded card it pairs
small round dot shapes with the text box beside them (and check icons with their sign-off text) and
merges them, keeping the spacing the slide already has.
"""
import argparse, copy, re
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS = {'a': A, 'p': P}
PX = 914400 / 144            # EMU per canvas px (1920 px = 13.333 in)
q = lambda tag: '{%s}%s' % (NS[tag.split(':')[0]], tag.split(':')[1])
TEAL, PULSE, MOSS, LEAF = '09A78D', '52E081', '0C2226', '194247'
DOT = ('●', 'Arial', 86)                 # char, font, size % of text
CHECK = ('ü', 'Wingdings', 110)          # Wingdings check mark, as in PowerPoint's own bullet menu
SEGOE_LINE = 1.33                             # PowerPoint single spacing for Segoe UI, as a multiple of size


# ---------------------------------------------------------------- paragraph building
def bullet_ppr(kind, color, marl_px, spc_bef_pt, ln=None):
    ppr = etree.Element(q('a:pPr'))
    m = int(round(marl_px * PX))
    ppr.set('marL', str(m)); ppr.set('indent', str(-m)); ppr.set('algn', 'l')
    if ln is not None:
        ppr.append(ln)
    if spc_bef_pt:
        sb = etree.SubElement(ppr, q('a:spcBef')); etree.SubElement(sb, q('a:spcPts')).set('val', str(int(round(spc_bef_pt * 100))))
    ch, font, pct = CHECK if kind == 'check' else DOT
    bc = etree.SubElement(ppr, q('a:buClr')); etree.SubElement(bc, q('a:srgbClr')).set('val', color)
    etree.SubElement(ppr, q('a:buSzPct')).set('val', str(pct * 1000))
    bf = etree.SubElement(ppr, q('a:buFont')); bf.set('typeface', font)
    if font == 'Wingdings':
        bf.set('pitchFamily', '2'); bf.set('charset', '2')
    etree.SubElement(ppr, q('a:buChar')).set('char', ch)
    return ppr


def cont_ppr(marl_px, ln=None):
    """Continuation paragraph inside one item (a line break in the source): same indent, no bullet."""
    ppr = etree.Element(q('a:pPr'))
    ppr.set('marL', str(int(round(marl_px * PX)))); ppr.set('indent', '0'); ppr.set('algn', 'l')
    if ln is not None:
        ppr.append(ln)
    etree.SubElement(ppr, q('a:buNone'))
    return ppr


def pct_spacing(lh_px, size_pt):
    ln = etree.Element(q('a:lnSpc'))
    pct = lh_px / (size_pt * 2 * SEGOE_LINE)
    etree.SubElement(ln, q('a:spcPct')).set('val', str(int(round(pct * 100)) * 1000))
    return ln


def text_paras(sp):
    """[(runs)] for each paragraph of a text shape (runs deep-copied, empty paragraphs skipped)."""
    out = []
    for p in sp.findall('.//p:txBody/a:p', NS):
        rs = [copy.deepcopy(r) for r in p if r.tag in (q('a:r'), q('a:br'), q('a:fld'))]
        if any(r.tag == q('a:r') and (r.findtext('a:t', namespaces=NS) or '').strip() for r in rs):
            out.append(rs)
    return out


def size_of(runs):
    for r in runs:
        rp = r.find('a:rPr', NS)
        if rp is not None and rp.get('sz'):
            return int(rp.get('sz')) / 100
    return 10.0


def make_box(tree, ref_el, x, y, w, h, paras, name):
    """paras: [(pPr, runs)]. Inserts the text box where ref_el is (keeps paint order and animation blocks)."""
    ids = [int(c.get('id')) for c in tree.iter(q('p:cNvPr')) if (c.get('id') or '').isdigit()]
    sp = etree.Element(q('p:sp'))
    nv = etree.SubElement(sp, q('p:nvSpPr'))
    c = etree.SubElement(nv, q('p:cNvPr')); c.set('id', str(max(ids + [1]) + 1)); c.set('name', name)
    etree.SubElement(nv, q('p:cNvSpPr')).set('txBox', '1')
    etree.SubElement(nv, q('p:nvPr'))
    spPr = etree.SubElement(sp, q('p:spPr'))
    xf = etree.SubElement(spPr, q('a:xfrm'))
    etree.SubElement(xf, q('a:off')).attrib.update({'x': str(int(x)), 'y': str(int(y))})
    etree.SubElement(xf, q('a:ext')).attrib.update({'cx': str(int(w)), 'cy': str(int(h))})
    g = etree.SubElement(spPr, q('a:prstGeom')); g.set('prst', 'rect'); etree.SubElement(g, q('a:avLst'))
    etree.SubElement(spPr, q('a:noFill'))
    tb = etree.SubElement(sp, q('p:txBody'))
    bp = etree.SubElement(tb, q('a:bodyPr'))
    bp.attrib.update({'wrap': 'square', 'lIns': '0', 'tIns': '0', 'rIns': '0', 'bIns': '0', 'rtlCol': '0', 'anchor': 't'})
    etree.SubElement(bp, q('a:spAutoFit'))
    etree.SubElement(tb, q('a:lstStyle'))
    for ppr, runs in paras:
        p = etree.SubElement(tb, q('a:p'))
        p.append(ppr)
        for r in runs:
            p.append(r)
        e = etree.SubElement(p, q('a:endParaRPr')); e.set('lang', 'en-US'); e.set('sz', str(int(size_of(runs) * 100)))
    ref_el.addprevious(sp)
    return sp


def geom(el):
    xf = el.find('.//p:spPr/a:xfrm', NS)
    if xf is None:
        return None
    o, e = xf.find('a:off', NS), xf.find('a:ext', NS)
    return int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))


# ---------------------------------------------------------------- tag mode (lists built with mkit)
TAG = re.compile(r'~(.*)$')


def split_tags(name):
    m = TAG.search(name or '')
    return (name[:m.start()], m.group(1).split('|')) if m else (name, [])


def convert_tagged(slide):
    tree = slide.shapes._spTree
    lists = {}
    for el in list(tree):
        c = el.find('.//p:cNvPr', NS)
        if c is None:
            continue
        base, parts = split_tags(c.get('name'))
        lp = next((p for p in parts if p.startswith('l:')), None)
        if not lp:
            continue
        lid, kind, color, gap, lh, idx, lx, ly, lw = lp[2:].split('/')
        rest = [p for p in parts if not p.startswith('l:')]
        L = lists.setdefault(lid, {'kind': kind, 'color': color, 'gap': float(gap), 'lh': float(lh), 'box': (int(lx), int(ly), int(lw)),
                                   'items': {}, 'rest': rest, 'first': el})
        L['items'].setdefault(int(idx), []).append(el)
    for lid, L in lists.items():
        lx, ly, lw = L['box']
        paras, marl, last_bottom = [], None, ly * PX
        for idx in sorted(L['items']):
            els = L['items'][idx]
            texts = [e for e in els if e.tag == q('p:sp') and e.find('.//p:cNvSpPr[@txBox="1"]', NS) is not None]
            if not texts:
                continue
            g = geom(texts[0])
            if marl is None:
                marl = (g[0] - lx * PX) / PX
            ps = [pp for t in texts for pp in text_paras(t)]
            size = size_of(ps[0]) if ps else 10
            for k, runs in enumerate(ps):
                ln = pct_spacing(L['lh'], size)
                if k == 0:
                    paras.append((bullet_ppr(L['kind'], L['color'], marl, (L['gap'] / 2) if paras else 0, ln), runs))
                else:
                    paras.append((cont_ppr(marl, ln), runs))
            last_bottom = max(last_bottom, max(geom(t)[1] + geom(t)[3] for t in texts))
        if not paras:
            continue
        name = 'List' + ('~' + '|'.join(L['rest']) if L['rest'] else '')
        make_box(tree, L['first'], lx * PX, ly * PX, lw * PX, max(last_bottom - ly * PX, 20 * PX), paras, name)
        for els in L['items'].values():
            for e in els:
                if e.getparent() is not None:
                    e.getparent().remove(e)
    # strip any leftover list tags from names
    for c in tree.iter(q('p:cNvPr')):
        base, parts = split_tags(c.get('name'))
        if any(p.startswith('l:') for p in parts):
            rest = [p for p in parts if not p.startswith('l:')]
            c.set('name', base + ('~' + '|'.join(rest) if rest else ''))
    return len(lists)


# ---------------------------------------------------------------- card mode (existing or hand-edited slides)
def _fill(sh):
    f = sh._element.find('.//p:spPr/a:solidFill/a:srgbClr', NS)
    return f.get('val').upper() if f is not None else None


def _is_dot(sh):
    g = sh._element.find('.//a:prstGeom', NS)
    return (sh.shape_type == 1 and g is not None and g.get('prst') == 'ellipse'
            and 5 * PX <= sh.width <= 13 * PX and 5 * PX <= sh.height <= 13 * PX)


def _within(sh, cd, pad=4 * PX):
    return cd.left - pad <= sh.left <= cd.left + cd.width and cd.top - pad <= sh.top <= cd.top + cd.height


def convert_cards(slide):
    shapes = list(slide.shapes)
    tree = slide.shapes._spTree
    cards = [s for s in shapes if s.shape_type == 1 and s.width > 200 * PX and s.height > 60 * PX
             and s._element.find('.//a:prstGeom', NS) is not None and s._element.find('.//a:prstGeom', NS).get('prst') == 'roundRect']
    done = 0
    for cd in cards:
        dark = _fill(cd) in (MOSS, LEAF)
        texts = [s for s in shapes if s.shape_type == 17 and s.has_text_frame and s.text_frame.text.strip() and _within(s, cd)]
        dots = sorted([s for s in shapes if _is_dot(s) and _within(s, cd)], key=lambda s: s.top)
        items = []
        for d in dots:
            cy = d.top + d.height / 2
            t = min((t for t in texts if 0 < t.left - d.left < 30 * PX and t.top - 6 * PX <= cy <= t.top + t.height),
                    key=lambda t: t.left - d.left, default=None)
            if t is not None:
                items.append((d, t))
        if items:
            x, y = items[0][0].left, items[0][1].top
            marl = (items[0][1].left - x) / PX
            gaps = [max(0, b.top - (a.top + a.height)) for (_, a), (_, b) in zip(items, items[1:])]
            gap_pt = max(3, round(sum(gaps) / len(gaps) / PX / 2)) if gaps else 4
            paras = []
            for k, (d, t) in enumerate(items):
                p0 = t._element.find('.//a:p', NS)
                ln = p0.find('a:pPr/a:lnSpc', NS)
                runs = text_paras(t._element)
                color = _fill(d) or (PULSE if dark else TEAL)
                for j, rs in enumerate(runs):
                    lnc = copy.deepcopy(ln) if ln is not None else None
                    paras.append((bullet_ppr('dot', color, marl, gap_pt if (k and not j) else 0, lnc) if j == 0 else cont_ppr(marl, lnc), rs))
            w = min(max(t.left + t.width for _, t in items), cd.left + cd.width - 20 * PX) - x
            h = items[-1][1].top + items[-1][1].height - y + 6 * PX
            make_box(tree, items[0][0]._element, x, y, w, h, paras, 'List')
            for d, t in items:
                for s in (d, t):
                    s._element.getparent().remove(s._element)
            done += 1
        # sign-off lines: a small check icon picture + a text box on the same line
        for pic in [s for s in shapes if s.shape_type == 13 and s.width < 30 * PX and _within(s, cd)]:
            if pic._element.getparent() is None:
                continue
            cy = pic.top + pic.height / 2
            t = next((t for t in texts if t._element.getparent() is not None and 0 < t.left - pic.left < 44 * PX
                      and t.top <= cy <= t.top + t.height), None)
            if t is None:
                continue
            marl = (t.left - pic.left) / PX
            paras = [(bullet_ppr('check', TEAL, marl, 0), rs) for rs in text_paras(t._element)[:1]]
            make_box(tree, pic._element, pic.left, t.top - 2 * PX, t.left + t.width - pic.left + 20 * PX, t.height + 4 * PX, paras, 'Sign-off')
            for s in (pic, t):
                s._element.getparent().remove(s._element)
    return done


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('src'); ap.add_argument('out')
    ap.add_argument('--cards', default=None, help='slide numbers to convert in card mode (existing slides), e.g. 2,5; "all" for every slide')
    a = ap.parse_args()
    prs = Presentation(a.src)
    if a.cards:
        want = set(range(1, len(prs.slides) + 1)) if a.cards == 'all' else {int(v) for v in a.cards.split(',')}
        for i, s in enumerate(prs.slides, 1):
            if i in want:
                print(f'slide {i}: {convert_cards(s)} lists merged (card mode)')
    else:
        n = sum(convert_tagged(s) for s in prs.slides)
        print(f'native bullets: {n} lists')
    prs.save(a.out)
